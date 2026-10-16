from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timezone

from openpyxl import load_workbook
import pytest

from core.domain.detail import DetailSection
from core.domain.values import FieldBag, NamedValue, PersonRef
from core.jira.audit import (
    JiraAuditUseCase,
    audit_issue,
    export_audit_xlsx,
    resolve_audit_input,
)
from core.jira.domain import Issue, IssueIdentity, JiraProjectRef, RichText
from core.ai import AIConfigurationError


@pytest.fixture(autouse=True)
def no_external_ai(monkeypatch):
    monkeypatch.setattr(
        "core.jira.audit.ai_review.create_chat_client",
        lambda _model: (_ for _ in ()).throw(AIConfigurationError("not configured")),
    )


GOOD_DESCRIPTION = """Steps to reproduce: Open video
Actual results: Freeze
Expected results: Playback
Reproducibility rate: 1/2
Comparison: Previous build works
Notes:
HW info: board A
SW info: build 1
"""


def _issue(key: str, creator: str, description: str = GOOD_DESCRIPTION) -> Issue:
    return Issue(
        IssueIdentity(key, key, f"https://jira.example/browse/{key}"),
        "[ACME][T7][V1][Video]: freezes", JiraProjectRef("SH"),
        NamedValue("1", "Open"), NamedValue("2", "Bug"),
        creator=PersonRef(creator.casefold(), creator.casefold(), creator),
        components=(NamedValue("10", "Video"),),
        description=DetailSection.loaded(RichText(description)),
        custom_fields=DetailSection.loaded(FieldBag.from_mapping({'Channel of Reporter': {'value': 'Customer-Feedback'}, 'Test Case ID': None})),
    )


@pytest.mark.parametrize('channel,test_case,failed', [
    ('Self-Test', None, True), ('Self-Test', '', True), ('Self-Test', '  ', True),
    ('Self-Test', 'TV-CATEGORY-CASE', False), ('Customer-Feedback', None, False), ('', None, False),
])
def test_self_test_requires_test_case_id_only_for_self_test(channel, test_case, failed):
    issue = replace(_issue('SH-1', 'Chao Li'), custom_fields=DetailSection.loaded(FieldBag.from_mapping(
        {'Channel of Reporter': {'value': channel}, 'Test Case ID': test_case})))
    result = audit_issue(issue)
    findings = [item for item in result.violations if item.rule_id == 'TEST_CASE_ID.REQUIRED']
    assert bool(findings) is failed
    assert result.passed is not failed
    if failed:
        assert findings[0].field == 'Test Case ID'
        assert 'Test Case ID 为空' in findings[0].reason


@pytest.mark.parametrize('value,passed', [
    ('TV-VIDEO-PLAY', True), ('IPTV_VIDEO_PLAY', True), ('OTT-video-play', True),
    ('NONE', True), ('None', True), ('CASE-1', False), ('TV-VIDEO_PLAY', False),
    ('TV--PLAY', False), ('TV-VIDEO-PLAY-EXTRA', False), ('TV-VIDEO-123', True),
    ('IPTV_Camera_012', True), ('IPTV-Camera-012', True), ('TV2_VIDEO2_CASE012', True),
    ('none', False),
])
def test_self_test_case_id_accepts_three_alphanumeric_segments_or_none(value, passed):
    issue = replace(_issue('SH-1', 'Chao Li'), custom_fields=DetailSection.loaded(FieldBag.from_mapping(
        {'Channel of Reporter': {'value': 'Self-Test'}, 'Test Case ID': value})))
    result = audit_issue(issue)
    assert result.passed is passed
    if not passed:
        assert result.violations[0].rule_id == 'TEST_CASE_ID.REQUIRED'
        assert result.violations[0].observed == value


def test_other_channels_do_not_require_test_case_field_mapping():
    issue = replace(_issue('SH-1', 'Chao Li'), custom_fields=DetailSection.loaded(FieldBag.from_mapping(
        {'Channel of Reporter': {'value': 'Customer-Feedback'}})))
    assert audit_issue(issue).passed
    issue = replace(issue, custom_fields=DetailSection.loaded(FieldBag.from_mapping(
        {'Channel of Reporter': {'value': 'Self-Test'}})))
    result = audit_issue(issue)
    assert not result.passed
    assert [item.rule_id for item in result.violations] == ['TEST_CASE_ID.REQUIRED']
    assert result.violations[0].reason == '当前 Jira 未提供 Test Case ID 参数。'
    assert audit_issue(replace(issue, custom_fields=DetailSection())).passed


@pytest.mark.parametrize('fields,reason', [
    ({}, '当前 Jira 未提供 Test Case ID 参数。'),
    ({'Test Case ID': None}, 'Test Case ID 为空。'),
    ({'Test Case ID': 'invalid'}, 'Test Case ID 格式不符合要求'),
])
def test_self_test_missing_case_id_uses_existing_export_and_remediation(tmp_path, fields, reason):
    from core.email.audit_report import summarize_audit, compare_jira_audits
    from core.jira.audit.models import AuditReport, JiraAuditScope
    issue = replace(_issue('SH-1', 'Chao Li'), custom_fields=DetailSection.loaded(FieldBag.from_mapping(
        {'Channel of Reporter': {'value': 'Self-Test'}, **fields})))
    from core.jira.audit.rules import active_rules
    report = AuditReport(JiraAuditScope('jql', '', 'key=SH-1'), datetime.now(timezone.utc), active_rules(), (audit_issue(issue),))
    summary = summarize_audit('jira', report)
    assert summary['failed'] == 1
    assert summary['findings'][0]['ruleId'] == 'TEST_CASE_ID.REQUIRED'
    assert reason in summary['findings'][0]['reason']
    path = export_audit_xlsx(report, output_path=tmp_path / 'jira.xlsx')
    workbook = load_workbook(path)
    assert any('TEST_CASE_ID.REQUIRED' in row and any(reason in str(value) for value in row)
               for row in workbook['违规明细'].values)
    fixed = replace(issue, custom_fields=DetailSection.loaded(FieldBag.from_mapping(
        {'Channel of Reporter': {'value': 'Self-Test'}, 'Test Case ID': 'TV-CATEGORY-CASE'})))
    rerun = replace(report, issues=(audit_issue(fixed),))
    assert compare_jira_audits(summary['findings'], rerun)[0]['state'] == 'fixed'


def test_jira_input_resolves_issue_search_and_filter_urls_and_validates_jql() -> None:
    validated = []
    fetch_filter = lambda filter_id: {"jql": f"filter = {filter_id}"}
    validate = lambda jql: validated.append(jql)

    issue = resolve_audit_input(
        "https://jira.example/browse/SH-10", base_url="https://jira.example",
        fetch_filter=fetch_filter, validate_jql=validate,
    )
    search = resolve_audit_input(
        "https://jira.example/issues/?jql=project%20%3D%20SH",
        base_url="https://jira.example", fetch_filter=fetch_filter,
        validate_jql=validate,
    )
    filtered = resolve_audit_input(
        "https://jira.example/filter/42", base_url="https://jira.example",
        fetch_filter=fetch_filter, validate_jql=validate,
    )
    filter_id = resolve_audit_input(
        "43", base_url="https://jira.example",
        fetch_filter=fetch_filter, validate_jql=validate,
    )

    assert (issue.source_kind, issue.jql) == ("issue_url", 'key = "SH-10"')
    assert (search.source_kind, search.jql) == ("jql_url", "project = SH")
    assert (filtered.source_kind, filtered.jql) == ("filter_url", "filter = 42")
    assert (filter_id.source_kind, filter_id.jql) == ("filter_id", "filter = 43")
    assert validated == ['key = "SH-10"', "project = SH", "filter = 42", "filter = 43"]
    with pytest.raises(ValueError, match="host"):
        resolve_audit_input(
            "https://other.example/browse/SH-1", base_url="https://jira.example",
            fetch_filter=fetch_filter, validate_jql=validate,
        )


def test_jira_deterministic_rules_keep_old_table_and_standard_behaviour() -> None:
    assert audit_issue(_issue("SH-1", "Chao Li")).passed
    table = "\n".join((
        "||||模块|||需要填写信息||||测试信息||||",
        "|平台信息|客户/项目代号|T6X高刷|", "|测试环境|测试仪器|-|",
    ))
    assert audit_issue(_issue("SH-2", "Chao Li", table)).passed

    failed = audit_issue(replace(
        _issue("SH-3", "Chao Li", ""),
        summary="bad", components=(),
    ))
    assert {item.rule_id for item in failed.violations} >= {
        "SUMMARY.FORMAT", "COMPONENT.REQUIRED",
        "DESCRIPTION.STEPS_TO_REPRODUCE", "DESCRIPTION.NOTES_HW",
    }


def test_description_rules_match_field_text_without_markup_requirements() -> None:
    description = """h2. [Steps to reproduce]
Open the application.
h2. [Actual results]:
The first screen takes more than 5 seconds.
h2. [Expected results]:
The first screen should load within 5 seconds.
h2. [Reproducibility rate]:
100%
h2. [Comparison]:
No comparison build.
h2. [Notes]
_HW info: board A_
SW info：[build 1]
"""

    result = audit_issue(_issue("IPTV-43297", "Chao Li", description))

    assert result.passed, {item.rule_id for item in result.violations}


def test_jira_use_case_loads_audit_details_only_for_eligible_creator() -> None:
    eligible = replace(_issue("SH-1", "Chao Li"), description=DetailSection())
    ignored = replace(_issue("SH-2", "Outside User"), description=DetailSection())

    class Source:
        def __init__(self): self.loaded = []
        def list_issues(self, _scope, _cancellation): return eligible, ignored
        def load_details(self, issue, details):
            self.loaded.append((issue.identity.key, details.sections()))
            return replace(issue, description=DetailSection.loaded(RichText(GOOD_DESCRIPTION)))

    source = Source()
    report = JiraAuditUseCase(source).run(
        type("Scope", (), {"source_kind": "jql", "original": "x", "jql": "x"})(),
    )

    assert [item.key for item in report.issues] == ["SH-1"]
    assert source.loaded == [("SH-1", ("description", "custom_fields"))]


def test_jira_exporter_creates_summary_and_violation_workbook(tmp_path) -> None:
    issue = replace(_issue("SH-1", "Chao Li", ""), components=())

    class Source:
        def list_issues(self, _scope, _cancellation): return (issue,)
        def load_details(self, issue, _details): return issue

    report = JiraAuditUseCase(Source()).run(
        type("Scope", (), {"source_kind": "jql", "original": "x", "jql": "project=SH"})(),
    )
    path = export_audit_xlsx(
        report, output_path=tmp_path / "jira.xlsx",
        now=datetime(2026, 8, 29, tzinfo=timezone.utc),
    )
    workbook = load_workbook(path)

    assert workbook.sheetnames == ["汇总", "违规明细"]
    assert tuple(workbook["汇总"].values)[8] == ("创建人", "违规数/创建数", "违规 Jira 号")
    detail = workbook['违规明细']
    assert detail['B1'].value == '人员（Creator）'
    assert detail['B2'].value == 'Chao Li'


def test_jira_use_case_starts_every_detail_when_requests_are_below_capacity() -> None:
    from threading import Barrier
    from core.async_tasks import AsyncTaskManager

    first = replace(_issue("SH-1", "Chao Li"), description=DetailSection())
    second = replace(_issue("SH-2", "Chao Li"), description=DetailSection())
    barrier = Barrier(2, timeout=1)

    class Source:
        def list_issues(self, _scope, _cancellation): return first, second
        def load_details(self, issue, _details):
            barrier.wait()
            return replace(issue, description=DetailSection.loaded(RichText(GOOD_DESCRIPTION)))

    manager = AsyncTaskManager(max_workers=3)
    try:
        root = manager.submit_coordinator(
            "jira-review", lambda token, _progress: JiraAuditUseCase(Source()).run(
                type("Scope", (), {"source_kind": "jql", "original": "x", "jql": "x"})(),
                task_manager=manager, parent_task_id=token.task_id,
            ),
        )
        report = root.result(timeout=2)
        root_id = manager.task_id(root)
        assert [item.key for item in report.issues] == ["SH-1", "SH-2"]
        assert len([task for task in manager._tasks.values() if task.parent_id == root_id]) == 2
    finally:
        manager.close()


def test_jira_root_cancellation_cancels_scheduled_issue_children() -> None:
    from core.async_tasks import AsyncTaskManager, TaskCancelled

    issues = tuple(replace(_issue(f"SH-{index}", "Chao Li"), description=DetailSection()) for index in (1, 2))
    manager = AsyncTaskManager(max_workers=3)
    root = manager.register_long_running("jira-review")

    class Source:
        def list_issues(self, _scope, _cancellation): return issues
        def load_details(self, issue, _details):
            if issue.identity.key == "SH-1":
                manager.cancel(root)
            return replace(issue, description=DetailSection.loaded(RichText(GOOD_DESCRIPTION)))

    try:
        with pytest.raises(TaskCancelled):
            JiraAuditUseCase(Source()).run(
                type("Scope", (), {"source_kind": "jql", "original": "x", "jql": "x"})(),
                task_manager=manager, parent_task_id=root,
            )
        assert all(task.token._event.is_set() for task in manager._tasks.values() if task.root_id == root)
    finally:
        manager.close()


def test_jira_review_uses_all_three_available_detail_workers() -> None:
    from threading import Barrier
    from core.async_tasks import AsyncTaskManager

    issues = tuple(replace(_issue(f"SH-{index}", "Chao Li"), description=DetailSection()) for index in (1, 2, 3))
    barrier = Barrier(3, timeout=1)

    class Source:
        def list_issues(self, _scope, _cancellation): return issues
        def load_details(self, issue, _details):
            barrier.wait()
            return replace(issue, description=DetailSection.loaded(RichText(GOOD_DESCRIPTION)))

    manager = AsyncTaskManager(max_workers=3)
    try:
        root = manager.submit_coordinator(
            "jira-review", lambda token, _progress: JiraAuditUseCase(Source()).run(
                type("Scope", (), {"source_kind": "jql", "original": "x", "jql": "x"})(),
                task_manager=manager, parent_task_id=token.task_id,
            ),
        )
        report = root.result(timeout=2)
        assert [item.key for item in report.issues] == ["SH-1", "SH-2", "SH-3"]
    finally:
        manager.close()
