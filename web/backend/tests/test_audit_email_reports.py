from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
import pytest
from core.confluence.audit import previous_business_week
from core.confluence.audit.models import AuditBatch, AuditFinding, AuditStatus, ProjectAudit
from core.confluence.audit.rules import UPDATE_MATRIX_POINTS
from core.confluence.project import Project, ProjectIdentity, ProductSpaceRef, ConfluencePageRef
from core.jira.audit import input as jira_input
from core.weekly_audit import fixed_weekly_audit_scope

from core.email.audit_report import render_history_report, summarize_audit
from core.jira.audit.models import AuditReport, IssueAuditResult, JiraAuditScope


def test_saved_confluence_comparisons_survive_schema_gate_without_rewrite(
    database_factory,
):
    from core.product_lines import PRODUCT_LINES
    from smarttest_web.audit.email_history import AuditEmailHistory
    from smarttest_web.schema import verify_schema_revision

    database = database_factory()
    history = AuditEmailHistory(database)
    values = ([6, 280], [18, 160], [16, 176], [7, 112])
    summary = {line.name: value for line, value in zip(PRODUCT_LINES, values)}
    summary.update(
        counts={
            line.name: {
                "updated": value[0],
                "not_updated": value[1] - value[0],
                "invalid_format": 0,
                "failed": 0,
                "unknown": 0,
            }
            for line, value in zip(PRODUCT_LINES, values)
        },
        projects=[{"projectId": "P1", "productLine": PRODUCT_LINES[0].name}],
        scope="saved period",
    )
    old = history.create("qa", {})
    old['summary']['confluence'] = summary
    old["reports"]["confluence"] = {
        "state": "completed",
        "html": "saved original body",
        "subject": "saved subject",
        "sourceIds": [old["id"]],
        "attachments": ["original.xlsx"],
    }
    old["deliveries"] = {"confluence": {"state": "sent", "sentAt": "saved-send-time"}}
    history.save("qa", old)
    current = history.create('qa', {})
    current["summary"]["confluence"] = summary
    records = [
        {"id": row["id"], "label": row["label"], "summary": summary}
        for row in (current, old)
    ]
    current["reports"]["confluence"] = {
        **render_history_report("confluence", records, current=True),
        "state": "completed",
        "sourceIds": [current["id"], old["id"]],
        "attachments": [],
    }
    history.save("qa", current)
    verify_schema_revision(database)
    verify_schema_revision(database)
    restored = AuditEmailHistory(database)
    saved = restored.get("qa", old["id"])
    for line, value in zip(PRODUCT_LINES, values):
        assert saved["summary"]["confluence"][line.name] == value
    assert saved["deliveries"] == old["deliveries"]
    assert saved["reports"]["confluence"] == old["reports"]["confluence"]
    html = restored.get("qa", current["id"])["reports"]["confluence"]["html"]
    assert html.count("6 / 280") == 2 and "0 / 0" not in html
    assert restored.get("qa", old["id"]) == saved


@pytest.mark.parametrize('space,name', [('DOPL', 'China Operator'),
    ('SDPL', 'Smart Device'), ('TV', 'TV'),
    ('OOPL', 'Global Operator & STB')])
def test_confluence_real_space_keys_are_summarized_in_canonical_product_lines(space, name, tmp_path):
    period = previous_business_week()
    project = Project(ProjectIdentity('1', 'P1'), 'Project', ProductSpaceRef(space), ConfluencePageRef('1'))
    finding = AuditFinding('P1', 'page', UPDATE_MATRIX_POINTS[0].rule_id, AuditStatus.NOT_UPDATED, 'missing')
    summary = summarize_audit('confluence', AuditBatch('b', period, datetime.now(), (ProjectAudit(project, (finding,)),)))
    assert summary[name] == [0, 1]
    assert summary['counts'][name]['not_updated'] == 1
    assert summary['projects'][0]['projectId'] == 'P1'
    assert summary['projects'][0]['productLine'] == name
    assert summary['projects'][0]['rules'][0]['status'] == 'not_updated'
    from core.confluence.audit.exporter import export_audit_xlsx_by_product_line
    paths = export_audit_xlsx_by_product_line(AuditBatch('b', period, datetime.now(), (ProjectAudit(project, (finding,)),)), tmp_path)
    assert paths[0].name == name + '_b.xlsx'


def test_remediation_explanation_is_readable_in_html_and_real_xlsx(tmp_path):
    from core.email.audit_report import remediation_table
    from core.reporting.excel import write_xlsx_sections
    from openpyxl import load_workbook
    records = [{'id': 'r', 'label': 'now', 'summary':
                {'total': 0, 'passed': 0, 'failed': 0, 'rate': '—', 'issues': [], 'scope': 'key=SH-1'}}]
    for index, remediation in enumerate((
            {'state': 'no_baseline', 'reason': 'no_adjacent_report', 'findings': []},
            {'state': 'completed', 'findings': []},
            {'state': 'completed', 'findings': [{'resourceId': 'SH-1', 'ruleId': 'r1',
                'reason': '原问题', 'state': 'fixed', 'currentReason': '当前通过', 'url': 'https://jira/SH-1'}]})):
        explanation, rows = remediation_table(remediation)
        if remediation['state'] == 'completed':
            remediation.update(jql='key=SH-1', originalAt='2026-09-14', reviewedAt='2026-09-18',
                originalSummary={'total': index - 1, 'passed': 0, 'failed': index - 1, 'rate': '—'},
                currentSummary={'total': index - 1, 'passed': index - 1, 'failed': 0, 'rate': '—'})
        html = render_history_report('jira', records, current=True, remediation=remediation)['html']
        assert '<table' in html[html.index('上期问题整改复查'):]
        if index == 0:
            assert explanation in html
        else:
            assert '同一保存 JQL' in html
        assert len(rows) == 1
        path = write_xlsx_sections(tmp_path / f'{index}.xlsx', sheet_name='整改', sections=[
            {'group': ('上期问题整改复查', explanation), 'headers': ('资源', '规则', '原问题', '状态', '证据', '链接'), 'rows': rows}])
        workbook = load_workbook(path)
        if index < 2:
            assert workbook.active.cell(3, 1).value == explanation
        else:
            assert tuple(cell.value for cell in workbook.active[3]) == (
                'SH-1', 'r1', '原问题', '已修复', '当前通过', 'https://jira/SH-1')
            assert '原问题' not in html and '当前通过' not in html
        assert '0%' not in explanation


def test_failed_jira_summary_preserves_stable_original_rule_details():
    from core.jira.audit.models import AuditViolation
    violation = AuditViolation('rule-1', 'description', 'steps', '', 'missing', 'fill')
    issue = IssueAuditResult('SH-1', 'https://jira/SH-1', '', 'QA', False, (violation,))
    report = AuditReport(JiraAuditScope('jql', '', 'project=SH'), datetime.now(), (), (issue,))
    summary = summarize_audit('jira', report)
    assert summary['findings'] == [{'resourceId': 'SH-1', 'ruleId': 'rule-1',
                                   'reason': 'missing', 'url': 'https://jira/SH-1'}]
    assert summary['issues'][0]['violations'][0] == {
        'ruleId': 'rule-1', 'section': 'description', 'field': 'steps',
        'reason': 'missing', 'guidance': 'fill', 'observed': ''}


def test_remediation_baseline_is_adjacent_completed_jira_and_same_account(
    tmp_path, *, database_factory
):
    from smarttest_web.audit.email_history import AuditEmailHistory
    from smarttest_web.database import WebDatabase
    history = AuditEmailHistory(database_factory())
    prior = {'startDate': '2026-09-07T00:00:00+08:00', 'endDate': '2026-09-12T00:00:00+08:00'}
    old = history.create('qa', prior)
    old['reports']['jira'] = {'state': 'completed'}
    old['summary']['jira'] = {'findings': [{'resourceId': 'SH-1', 'ruleId': 'r1'}]}
    old['deliveries'] = {'jira': {'state': 'failed'}}
    history.save('qa', old)
    assert history.previous_jira_audit('qa', prior)['id'] == old['id']
    assert history.previous_jira_audit('other', prior) is None
    newest = history.create('qa', prior)
    newest['reports']['jira'] = {'state': 'completed'}
    newest['deliveries'] = old['deliveries']
    history.save('qa', newest)
    assert history.previous_jira_audit('qa', prior)['id'] == newest['id']


def test_jira_saved_filter_report_comparison_never_calls_missing_candidates_fixed():
    from core.email.audit_report import compare_jira_audits
    from core.jira.audit.models import AuditRule, AuditViolation
    original = [{'resourceId': key, 'ruleId': rule, 'reason': 'old'}
                for key, rule in [('SH-1', 'r1'), ('SH-2', 'r1'), ('SH-3', 'r1'), ('SH-1', 'removed')]]
    report = AuditReport(JiraAuditScope('jql', '', 'project=SH'), datetime.now(),
        (AuditRule('r1', '', '', '', ''),), (
            IssueAuditResult('SH-1', '', '', '', True, ()),
            IssueAuditResult('SH-2', '', '', '', False, (AuditViolation('r1', '', '', '', 'still missing', ''),))))
    results = compare_jira_audits(original, report)
    assert [item['state'] for item in results] == ['fixed', 'unfixed', 'unverifiable', 'unverifiable']
    assert '无法确认' in results[2]['currentReason']


def test_jira_body_shows_same_jql_two_audit_summaries_not_rule_details():
    normal = {'total': 10, 'passed': 9, 'failed': 1, 'rate': '90%', 'issues': [], 'scope': 'project=TV'}
    old = {'total': 2, 'passed': 0, 'failed': 2, 'rate': '0%'}
    new = {'total': 2, 'passed': 1, 'failed': 1, 'rate': '50%'}
    remediation = {'state': 'completed', 'jql': 'project=SH AND summary ~ "<text>"',
        'originalAt': '2026-09-14T07:00:00+00:00', 'reviewedAt': '2026-09-18T08:00:00+00:00',
        'originalSummary': old, 'currentSummary': new,
        'findings': [{'resourceId': 'SH-1', 'ruleId': 'private-rule', 'reason': 'private-detail', 'state': 'fixed'}]}
    html = render_history_report('jira', [{'id': 'normal', 'label': 'now', 'summary': normal}],
                                 current=True, remediation=remediation)['html']
    section = html[html.index('上期问题整改复查'):]
    assert 'project=SH AND summary ~ &quot;&lt;text&gt;&quot;' in section
    assert '50%' in section and '0%' in section and '90%' not in section
    assert '2026-09-14 15:00:00 北京时间' in section
    assert '2026-09-18 16:00:00 北京时间' in section
    assert 'private-rule' not in section and 'private-detail' not in section
    assert 'background:#fff2cc' in section and 'background:#dbeafe' in section


def test_report_counts_issues_once_and_escapes_content_and_unsafe_links():
    issue = IssueAuditResult('SH-1', 'javascript:alert(1)', '', '<script>owner</script>', False, ())
    report = AuditReport(JiraAuditScope('jql', '', 'project="<SH>"'), datetime.now(), (), (issue, issue))
    summary = summarize_audit('jira', report)
    assert summary['total'] == 1 and summary['failed'] == 1
    rendered = render_history_report('jira', [{'id': 'run', 'label': 'today', 'summary': summary}], current=True)['html']
    assert '<script>' not in rendered
    assert 'javascript:' not in rendered
    assert '&lt;script&gt;owner&lt;/script&gt;' in rendered
    assert 'SH-1' in rendered
    assert '下面是针对大家本周创建的bug进行的规范检查，针对还不满足规范的部分，大家需要尽快改善。' in rendered
    assert '本次已执行审查，报告已保存。' not in rendered
    assert '<h2>FAE QA JIRA描写不符合规范反馈</h2>' in rendered
    assert '<h2>FAE QA JIRA描写不符合规范反馈 today</h2>' not in rendered


def test_zero_issue_report_has_no_invented_pass_rate():
    summary = summarize_audit('jira', AuditReport(JiraAuditScope('jql', '', 'project=SH'), datetime.now(), (), ()))
    assert summary['rate'] == '—'
    assert summary['total'] == 0


@pytest.mark.parametrize('now,start,end', [(datetime(2026, 9, 7), '2026-08-31', '2026-09-07'),
                                         (datetime(2027, 1, 4), '2026-12-28', '2027-01-04')])
def test_saved_created_bounds_roll_with_business_week_without_changing_other_conditions(now, start, end):
    period = previous_business_week(now.replace(tzinfo=ZoneInfo('Asia/Shanghai')))
    template = 'project in (SH, TV, IPTV, OTT,RK) AND issuetype in (Bug, Sub-bug) AND created >=2026-08-28 AND created <=2026-09-4 order by updated DESC'
    expected = f'project in (SH, TV, IPTV, OTT,RK) AND issuetype in (Bug, Sub-bug) AND created >="{start}" AND created <"{end}" order by updated DESC'
    assert jira_input.weekly_audit_jql(template, period) == expected


def test_jql_without_created_comparisons_is_unchanged_including_quoted_text():
    query = 'project=SH AND summary ~ "created >=2026-08-28" ORDER BY created DESC'
    assert jira_input.weekly_audit_jql(query, previous_business_week()) == query


@pytest.mark.parametrize('trigger,start,end,commercial_years', [
    (datetime(2026, 9, 18, 18, 0, tzinfo=ZoneInfo('Asia/Shanghai')), '2026-09-11T18:00:00+08:00', '2026-09-18T18:00:00+08:00', ['2025', '2026']),
    (datetime(2026, 9, 16, 9, 30, tzinfo=ZoneInfo('Asia/Shanghai')), '2026-09-11T18:00:00+08:00', '2026-09-18T18:00:00+08:00', ['2025', '2026']),
    (datetime(2027, 1, 1, 18, 0, tzinfo=ZoneInfo('Asia/Shanghai')), '2026-12-25T18:00:00+08:00', '2027-01-01T18:00:00+08:00', ['2026', '2027']),
])
def test_fixed_weekly_scope_is_continuous_from_last_friday_18_to_this_friday_18(trigger, start, end, commercial_years):
    scope = fixed_weekly_audit_scope(trigger)
    assert scope['startDate'] == start
    assert scope['endDate'] == end
    previous_start = datetime.fromisoformat(start) - timedelta(days=7)
    assert scope['previousPeriod'] == {
        'startDate': previous_start.isoformat(),
        'endDate': start,
    }
    assert scope['jira']['jql'] == (
        'project in (IPTV, SH, TV, OTT,RK) AND issuetype in (Bug, Sub-bug) '
        f'AND created >= "{datetime.fromisoformat(start):%Y-%m-%d %H:%M}" '
        f'AND created < "{datetime.fromisoformat(end):%Y-%m-%d %H:%M}" order by updated DESC'
    )
    assert '+08:00' not in scope['jira']['jql']
    assert scope['confluence']['filters'] == {
        'product space': ['SDPL', 'TV', 'OOPL'],
        'date of commercial approval': commercial_years,
        'support mode': ['A', 'B'],
        'project status': ['NORMAL'],
    }
    assert scope['confluence']['excludeCurrentStageAtOrAbove'] == 4
    assert scope['confluence']['excludeSupportModeBProductLines'] == ['SDPL']


def test_confluence_denominator_counts_all_actual_update_point_statuses_by_product_line():
    period = previous_business_week()
    project = Project(ProjectIdentity('1', 'P1'), 'Project', ProductSpaceRef('TV'), ConfluencePageRef('1'))
    states = [AuditStatus.UPDATED, AuditStatus.NOT_UPDATED, AuditStatus.INVALID_FORMAT, AuditStatus.FAILED,
              AuditStatus.UNKNOWN, AuditStatus.UPDATED, AuditStatus.NOT_UPDATED, AuditStatus.NOT_UPDATED]
    findings = tuple(AuditFinding('P1', 'Test', point.rule_id, status, '')
                     for point, status in zip(UPDATE_MATRIX_POINTS, states))
    findings += (AuditFinding('P1', 'Test', 'role.major_fae_qa', AuditStatus.UNKNOWN, ''),)
    summary = summarize_audit('confluence', AuditBatch('batch', period, datetime.now(), (ProjectAudit(project, findings),)))
    assert summary['TV'] == [2, 8]
    assert summary['China Operator'] == [0, 0]
    html = render_history_report('confluence', [{'id': 'run', 'label': 'today', 'summary': summary}], current=True)['html']
    assert '2 / 8' in html
    assert '下面是本周confluence信息更新检查结果，请未更新的项目owner尽快去补充未完成的部分。' in html
    assert '<h2>Confluence信息更新检查结果</h2>' in html
    assert '<h2>Confluence信息更新检查结果 today</h2>' not in html
    assert '<th style="border:1px solid #bcc9da;padding:10px;text-align:left;background:#dbeafe">格式有误</th>' not in html
    assert '>失败</th>' in html and '>未知</th>' in html
    assert '待确认' not in html and '尚未确认' not in html
    summary['TV'][1] = None
    with pytest.raises(ValueError):
        render_history_report('confluence', [{'id': 'old', 'label': 'old', 'summary': summary}])


@pytest.mark.parametrize('damage', ['missing_line', 'missing_counts', 'inconsistent_counts'])
def test_confluence_report_rejects_incomplete_or_inconsistent_statistics(damage):
    from core.product_lines import PRODUCT_LINES
    summary = {line.name: [0, 0] for line in PRODUCT_LINES}
    if damage == 'missing_line':
        del summary['China Operator']
    else:
        summary['counts'] = {line.name: dict(updated=0, not_updated=0, invalid_format=0, failed=0, unknown=0)
                             for line in PRODUCT_LINES}
        if damage == 'missing_counts':
            del summary['counts']['China Operator']
        else:
            summary['counts']['China Operator']['updated'] = 1
    with pytest.raises(ValueError):
        render_history_report('confluence', [{'id': 'old', 'label': 'old', 'summary': summary}])


def test_weekly_creator_counts_are_deduplicated_identity_scoped_and_saved_with_period(
    tmp_path, *, database_factory
):
    from core.jira.audit.models import AuditViolation
    from smarttest_web.audit.email_history import AuditEmailHistory
    from smarttest_web.database import WebDatabase
    violation = AuditViolation('rule', 'section', 'field', '', 'bad', 'fix')
    issues = tuple(IssueAuditResult(f'SH-{index}', f'https://jira/SH-{index}', '', 'Same Name', index >= 2,
        (violation, violation) if index < 2 else (), creator_identity='stable-1', creator_account='account.one') for index in range(10))
    other = IssueAuditResult('SH-other', 'https://jira/SH-other', '', 'Same Name', False, (violation,), creator_identity='stable-2', creator_account='account.two')
    clean = IssueAuditResult('SH-clean', '', '', 'Zero Violations', True, (), creator_identity='stable-3', creator_account='account.three')
    report = AuditReport(JiraAuditScope('jql', '', 'saved created week'), datetime.now(), (), (*issues, issues[0], other, clean))
    period = {'startDate': '2026-10-02T18:00:00+08:00', 'endDate': '2026-10-09T18:00:00+08:00'}
    summary = summarize_audit('jira', report, period=period)
    assert summary['period'] == period
    assert summary['people'] == [
        {'identity': 'stable-1', 'account': 'account.one', 'displayName': 'Same Name', 'createdCount': 10, 'violatedCount': 2},
        {'identity': 'stable-2', 'account': 'account.two', 'displayName': 'Same Name', 'createdCount': 1, 'violatedCount': 1},
        {'identity': 'stable-3', 'account': 'account.three', 'displayName': 'Zero Violations', 'createdCount': 1, 'violatedCount': 0},
    ]
    assert summary['issues'][0]['creatorIdentity'] == 'stable-1'
    assert summary['issues'][0]['creatorAccount'] == 'account.one'
    history = AuditEmailHistory(database_factory())
    run = history.create('qa', period)
    history.complete_report('qa', run, 'jira', summary, [])
    saved = history.get('viewer', run['id'])
    assert saved['summary']['jira'] == summary
    html = saved['reports']['jira']['html']
    assert '违规数/创建数' in html and '>2/10<' in html and '>1/1<' in html
    assert 'Zero Violations' not in html

    from core.jira.audit.exporter import export_audit_xlsx
    from openpyxl import load_workbook
    sheet = load_workbook(export_audit_xlsx(report, output_path=tmp_path / 'review.xlsx'))['汇总']
    assert tuple(sheet.cell(9, column).value for column in range(1, 4)) == ('创建人', '违规数/创建数', '违规 Jira 号')
    assert sheet.cell(10, 2).value == '2/10'
    assert sheet.cell(10, 3).value == 'SH-0、SH-1'
    assert sheet.cell(11, 2).value == '1/1'
